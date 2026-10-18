import streamlit as st
import os
import json
from dotenv import load_dotenv
from langchain_groq import ChatGroq
from pypdf import PdfReader
from pptx import Presentation  

load_dotenv()
st.set_page_config(page_title="MI Tananyag Feldolgozó", layout="wide", page_icon="📚")

if "ce_data" not in st.session_state:
    st.session_state.ce_data = None

st.sidebar.header("⚙️ Beállítások")
api_key = os.getenv("GROQ_API_KEY")

if api_key:
    st.sidebar.success("✅ API kulcs aktív")
else:
    st.sidebar.error("❌ API kulcs hiányzik!")
    st.stop()

temp = st.sidebar.slider("Kreativitás", 0.0, 1.0, 0.4)
style = st.sidebar.selectbox("Stílus", ["Egyetemi", "Konyhanyelv", "Lényegre törő"])

st.title("📚 Intelligens Dokumentum Feldolgozó")

input_method = st.radio("Hogyan szeretnéd megadni a tananyagot?", ["Szöveg bemásolása", "Fájl feltöltése (PDF vagy PPTX)"])

final_text = ""

if input_method == "Szöveg bemásolása":
    final_text = st.text_area("Másold be a szöveget:", height=200)
else:
    uploaded_file = st.file_uploader("Válassz ki egy fájlt", type=["pdf", "pptx"]) 
    if uploaded_file is not None:
        file_ext = uploaded_file.name.split(".")[-1].lower()
        
        with st.spinner(f"{file_ext.upper()} feldolgozása..."):
            if file_ext == "pdf":
                reader = PdfReader(uploaded_file)
                for page in reader.pages:
                    final_text += page.extract_text() + "\n"
            
            elif file_ext == "pptx":
                prs = Presentation(uploaded_file)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            final_text += shape.text + "\n"
            
            st.success(f"Sikeresen beolvasva: {len(final_text)} karakter.")

def build_prompt(text, s):
    return f"""
    Te egy oktatási szakértő vagy. Elemezd a forrásszöveget.
    STÍLUS: {s}
    FORRÁSSZÖVEG: {text[:8000]}
    
    VÁLASZ: Kizárólag érvényes JSON formátumban válaszolj:
    {{
      "summary": "hosszú összefoglaló",
      "keywords": [{{"term": "szó", "definition": "leírás"}}],
      "quiz": [{{"question": "kérdés", "options": ["A", "B", "C"], "answer": "helyes válasz szövege"}}]
    }}
    """

if st.button("✨ Elemzés indítása"):
    if not final_text or len(final_text) < 50:
        st.warning("Nincs elég szöveg a feldolgozáshoz!")
    else:
        try:
            with st.spinner("Az MI elemzi a dokumentumot..."):
                model = ChatGroq(model="llama-3.3-70b-versatile", temperature=temp, groq_api_key=api_key)
                response = model.invoke(build_prompt(final_text, style))
                
                content = response.content.strip()
                if "```json" in content:
                    content = content.split("```json")[1].split("```")[0].strip()
                elif "```" in content:
                    content = content.split("```")[1].split("```")[0].strip()
                
                st.session_state.ce_data = json.loads(content)
                st.rerun()
        except Exception as e:
            st.error("Hiba történt az elemzés során.")
            st.exception(e)

if st.session_state.ce_data:
    data = st.session_state.ce_data
    st.divider()
    t1, t2, t3 = st.tabs(["📖 Összefoglaló", "🔑 Fogalmak", "🧠 Kvíz"])

    with t1:
        st.write(data["summary"])
    with t2:
        for item in data["keywords"]:
            with st.expander(item["term"]):
                st.write(item["definition"])
    with t3:
        for i, q in enumerate(data["quiz"]):
            st.markdown(f"**{i+1}. {q['question']}**")
            choice = st.radio(f"Válasz ({i+1}):", q["options"], key=f"q_{i}", index=None)
            if choice == q["answer"]:
                st.success("Helyes!")
            elif choice:
                st.info(f"Próbáld újra! A helyes: {q['answer']}")

    if st.button("🗑️ Reset"):
        st.session_state.ce_data = None
        st.rerun()